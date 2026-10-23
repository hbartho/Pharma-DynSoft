"""Import d'approvisionnements par fichier CSV / Excel : modèle téléchargeable + aperçu (lignes reconnues / en erreur).
La création de l'appro se fait ensuite via POST /supplies avec les lignes reconnues.
"""
import io
import re
import csv
import unicodedata
from datetime import datetime, date
from typing import Optional, List, Dict, Any

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import StreamingResponse

from auth import get_current_user
from database.config import db_manager
from database.tenant_context import get_current_tenant

router = APIRouter(prefix="/supplies/import", tags=["Supplies Import"])

TEMPLATE_COLUMNS = [
    ("reference", "Référence"), ("barcode", "Code-barres"), ("product", "Produit"),
    ("quantity", "Quantité"), ("unit_price", "Prix cession HT"), ("selling_price", "Prix public"),
    ("expiration", "Date péremption"), ("lot", "N° lot"), ("tva", "TVA %"),
]
TEMPLATE_EXAMPLES = [
    ["", "", "Paracétamol 500mg", 20, 1000, 1500, "2027-12-31", "LOT-PAR-2411", 0],
    ["", "6181234567890", "", 10, 2500, 3500, "31/08/2027", "", 0],
    ["REF-AMOX1G", "", "", 30, 1800, "", "2028-03-15", "AMX-0325", 0],
    ["", "", "Ibuprofène 400mg", 12, 900, 1300, "30/06/2027", "", 0],
    ["", "", "Ibuprofène 400mg", 8, 900, 1300, "2027-11-30", "IBU-B2", 0],
]

# Alias de colonnes acceptés (comparés après normalisation : minuscules, sans accents, sans espaces)
ALIASES = {
    "reference": ["reference", "ref", "refinterne", "referenceinterne", "code", "codeproduit"],
    "barcode": ["codebarres", "codebarre", "barcode", "ean", "ean13", "cb"],
    "product": ["produit", "product", "nom", "designation", "libelle", "nomproduit", "article"],
    "quantity": ["quantite", "qte", "qty", "quantity", "nombre"],
    "unit_price": ["prixcessionht", "prixcession", "prixachat", "pu", "puht", "prixunitaire", "unitprice", "purchaseprice", "cession", "prixappro"],
    "selling_price": ["prixpublic", "prixvente", "prixpublicttc", "pv", "sellingprice", "ttc", "prixventettc"],
    "expiration": ["dateperemption", "peremption", "expiration", "expirationdate", "dlc", "dluo", "dateexpiration", "exp"],
    "lot": ["nlot", "lot", "numerolot", "numlot", "lotnumber", "nolot"],
    "tva": ["tva", "tva%", "tauxtva", "vat", "tax", "taxrate"],
}


def norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9%]", "", s.lower())


def map_columns(headers: List[str]) -> Dict[str, int]:
    mapping = {}
    normalized = [norm(h) for h in headers]
    for key, aliases in ALIASES.items():
        for idx, h in enumerate(normalized):
            if h in aliases and key not in mapping:
                mapping[key] = idx
    return mapping


def to_float(v) -> Optional[float]:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v)).replace(" ", "")
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    elif "," in s and "." in s:
        s = s.replace(",", "")
    return float(s) if s not in ("", "-", ".") else None


def to_date(v) -> Optional[date]:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d", "%m/%Y", "%Y-%m"):
        try:
            d = datetime.strptime(s, fmt).date()
            return d.replace(day=28) if fmt in ("%m/%Y", "%Y-%m") else d
        except ValueError:
            continue
    return None


def read_rows(filename: str, content: bytes) -> List[List[Any]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        return [[c for c in row] for row in ws.iter_rows(values_only=True)]
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(status_code=400, detail="Encodage du fichier non reconnu")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), dialect)]


@router.get("/template")
def download_template(format: str = Query(default="xlsx"), current_user: dict = Depends(get_current_user)):
    headers = [label for _, label in TEMPLATE_COLUMNS]
    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        w.writerow(headers)
        w.writerows(TEMPLATE_EXAMPLES)
        data = ("\ufeff" + buf.getvalue()).encode("utf-8")
        return StreamingResponse(io.BytesIO(data), media_type="text/csv; charset=utf-8",
                                 headers={"Content-Disposition": 'attachment; filename="modele_approvisionnement.csv"'})
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Approvisionnement"
    ws.append(headers)
    for row in TEMPLATE_EXAMPLES:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for col, width in zip("ABCDEFGHI", (14, 16, 34, 10, 16, 14, 16, 14, 8)):
        ws.column_dimensions[col].width = width
    notes = wb.create_sheet("Aide")
    for line in [
        "Une ligne par produit. Colonnes obligatoires : Produit (ou Référence / Code-barres), Quantité, Prix cession HT, Date péremption.",
        "Le produit est reconnu par Référence, puis Code-barres, puis Nom exact. Les produits inconnus sont signalés et ignorés.",
        "Prix public vide = inchangé. N° lot vide = repère automatique. Dates : AAAA-MM-JJ ou JJ/MM/AAAA.",
        "Même produit avec plusieurs dates de péremption = une ligne par lot (voir Ibuprofène dans l'exemple).",
        "Les lignes d'exemple sont à remplacer par vos produits : elles ne correspondent pas forcément à votre catalogue.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="modele_approvisionnement.xlsx"'})


@router.post("/preview")
def preview_import(file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    """Analyse le fichier et retourne les lignes prêtes (product_id résolu) et celles en erreur."""
    from database.models_tenant import Product

    content = file.file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Fichier trop volumineux (5 Mo max)")
    rows = read_rows(file.filename, content)
    rows = [r for r in rows if r and any(c not in (None, "") for c in r)]
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="Fichier vide : il faut une ligne d'en-tête et au moins une ligne de produit")
    mapping = map_columns([str(c or "") for c in rows[0]])
    missing = [k for k in ("quantity", "unit_price") if k not in mapping]
    if not any(k in mapping for k in ("reference", "barcode", "product")):
        missing.insert(0, "product")
    if missing:
        labels = dict(TEMPLATE_COLUMNS)
        raise HTTPException(status_code=400, detail=f"Colonnes manquantes : {', '.join(labels[m] for m in missing)}. Téléchargez le modèle pour voir le format attendu.")

    with db_manager.get_tenant_session(get_current_tenant()) as session:
        products = session.query(Product).filter(Product.is_active == True).all()  # noqa: E712
        by_ref = {norm(p.internal_reference): p for p in products if p.internal_reference}
        by_barcode = {norm(p.barcode): p for p in products if p.barcode}
        by_name = {norm(p.name): p for p in products if p.name}
        by_name_raw = {p.name.strip().lower(): p for p in products if p.name}

        def get(row, key):
            idx = mapping.get(key)
            return row[idx] if idx is not None and idx < len(row) else None

        result = []
        for line_no, row in enumerate(rows[1:], start=2):
            ref, barcode, name = get(row, "reference"), get(row, "barcode"), get(row, "product")
            product = None
            matched_by = None
            if ref and norm(ref) in by_ref:
                product, matched_by = by_ref[norm(ref)], "référence"
            elif barcode and norm(barcode) in by_barcode:
                product, matched_by = by_barcode[norm(barcode)], "code-barres"
            elif name and (str(name).strip().lower() in by_name_raw or norm(name) in by_name):
                product = by_name_raw.get(str(name).strip().lower()) or by_name[norm(name)]
                matched_by = "nom"
            qty = to_float(get(row, "quantity"))
            unit_price = to_float(get(row, "unit_price"))
            selling_price = to_float(get(row, "selling_price"))
            exp = to_date(get(row, "expiration"))
            tva = to_float(get(row, "tva")) or 0
            lot = str(get(row, "lot") or "").strip()
            errors = []
            if not product:
                errors.append(f"Produit introuvable ({ref or barcode or name or 'vide'})")
            if qty is None or qty <= 0 or int(qty) != qty:
                errors.append("Quantité invalide")
            if unit_price is None or unit_price < 0:
                errors.append("Prix cession invalide")
            if exp is None:
                errors.append("Date de péremption manquante ou invalide")
            elif exp <= date.today():
                errors.append("Date de péremption déjà passée")
            result.append({
                "line": line_no,
                "status": "ok" if not errors else "error",
                "errors": errors,
                "matched_by": matched_by,
                "product_id": str(product.id) if product else None,
                "product_name": product.name if product else (str(name or ref or barcode or "").strip() or None),
                "current_price": product.price if product else None,
                "quantity": int(qty) if qty and qty > 0 else None,
                "unit_price": unit_price,
                "selling_price": selling_price,
                "date_peremption": exp.isoformat() if exp else None,
                "lot_number": lot or None,
                "tva_rate": tva,
            })

    ok = [r for r in result if r["status"] == "ok"]
    return {
        "filename": file.filename,
        "columns_detected": list(mapping.keys()),
        "rows": result,
        "summary": {
            "total": len(result), "ok": len(ok), "errors": len(result) - len(ok),
            "total_amount": round(sum(r["quantity"] * r["unit_price"] for r in ok), 2),
        },
    }
